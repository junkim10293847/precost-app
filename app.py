"""사전원가 자동입력 — Streamlit 앱
실행: streamlit run app.py
로그인 계정은 Streamlit Cloud → App settings → Secrets 에 등록 (README 참고)
"""
import io, os, hmac, hashlib, tempfile, urllib.parse
import pandas as pd
import streamlit as st
import openpyxl

from core import (read_item_list, collect_files, run, to_excel, read_owner_master, assign_owners,
                  owner_requests, request_mail, read_postcost, compare_pre_post, reselect_file, refresh_review)

st.set_page_config(page_title='사전원가 자동입력', page_icon='📑', layout='wide')


# ================================================================ 로그인
def _users():
    try:
        return dict(st.secrets.get('users', {}))
    except Exception:
        return {}


def _check(pw_input, stored):
    stored = str(stored)
    if stored.startswith('sha256:'):
        return hmac.compare_digest(hashlib.sha256(pw_input.encode()).hexdigest(), stored[7:])
    return hmac.compare_digest(pw_input, stored)


def login_gate():
    if st.session_state.get('user'):
        return
    users = _users()
    st.title('📑 사전원가 자동입력')
    if not users:
        st.error('로그인 계정이 아직 등록되지 않았습니다. 관리자: Streamlit Cloud → Manage app → Settings → '
                 'Secrets 에 [users] 를 등록하세요.')
        st.code('[users]\nadmin = "비밀번호"\nfinance1 = "비밀번호"', language='toml')
        st.stop()
    with st.form('login'):
        uid = st.text_input('아이디')
        pw = st.text_input('비밀번호', type='password')
        ok = st.form_submit_button('로그인', type='primary')
    if ok:
        if uid in users and _check(pw, users[uid]):
            st.session_state['user'] = uid
            st.rerun()
        else:
            st.error('아이디 또는 비밀번호가 맞지 않습니다.')
    st.stop()


login_gate()

# ================================================================ 화면
st.title('📑 사전원가 자동입력')
st.caption('요청 품목 목록 + 견적서(zip) → 원가 자동입력 · 화면에서 수정/확인 후 저장 · 담당자별 요청 메일 · 사전/사후원가 비교')

with st.sidebar:
    st.markdown(f"👤 **{st.session_state['user']}** 님")
    if st.button('로그아웃'):
        for k in list(st.session_state.keys()):
            del st.session_state[k]
        st.rerun()
    st.divider()
    st.header('사용 방법')
    st.markdown(
        '1. **요청 품목 목록** 엑셀\n2. **견적서** zip 또는 여러 파일\n'
        '3. (선택) **지난달 결과 파일** → 담당자 자동 배정\n4. (선택) **사후원가 파일** → 사전/사후 비교\n'
        '5. 실행 → 노란 행 수정·**확인 체크** → **💾 저장** → 다운로드\n\n'
        '**파일명 규칙**: 견적서 파일명 앞에 목록 번호 (예: `29. 테아닌.xlsx`, `82~85.`)')
    period = st.text_input('대상 월', value='2026년 7월')
    sender = st.text_input('발신 부서', value='재무지원팀')
    due = st.text_input('회신 기한(선택)', value='')
    match_mode = st.radio('견적서 매칭 방식', ['파일번호 + 품목명 (권장)', '품목명만 (번호 없는 파일)'], index=0)

c1, c2 = st.columns(2)
with c1:
    list_file = st.file_uploader('① 사전원가 요청 품목 목록 (엑셀)', type=['xlsx', 'xlsm'])
    sheet = None
    if list_file:
        names = openpyxl.load_workbook(io.BytesIO(list_file.getvalue()), read_only=True).sheetnames
        default = next((i for i, n in enumerate(names) if '사전원가' in n or 'List' in n), 0)
        sheet = st.selectbox('목록 시트', names, index=default)
    owner_files = st.file_uploader('③ (선택) 지난달 결과 파일 — 담당자 자동 배정용', type=['xlsx'],
                                   accept_multiple_files=True)
with c2:
    quote_files = st.file_uploader('② 견적서 파일 (zip 또는 여러 파일)', accept_multiple_files=True,
                                   type=['zip', 'xlsx', 'xls', 'xlsm', 'pdf', 'png', 'jpg', 'jpeg'])
    post_file = st.file_uploader('④ (선택) 사후원가 파일 — 제품명·제조번호·원료비·부자재비·노무비/경비·생산량',
                                 type=['xlsx'])
    post_sheet, post_unit = None, None
    if post_file:
        pn = openpyxl.load_workbook(io.BytesIO(post_file.getvalue()), read_only=True).sheetnames
        post_sheet = st.selectbox('사후원가 시트', pn)
        opt = st.radio('사후원가 금액 형태', ['자동(생산량 열 있으면 금액÷생산량)', '이미 단위원가', '금액 → 생산량으로 나눔'],
                       horizontal=True)
        post_unit = {'자동(생산량 열 있으면 금액÷생산량)': None, '이미 단위원가': True, '금액 → 생산량으로 나눔': False}[opt]

if st.button('🚀 자동 입력 실행', type='primary', disabled=not (list_file and quote_files)):
    try:
        items, _ = read_item_list(list_file.getvalue(), sheet)
    except Exception as e:
        st.error(f'목록을 읽지 못했습니다: {e}')
        st.stop()
    files = collect_files([(f.name, f.getvalue()) for f in quote_files], tempfile.mkdtemp())
    bar = st.progress(0.0, text='준비 중…')
    rows = run(items, files, progress=lambda i, n, nm: bar.progress(i / n, text=f'{i}/{n} · {nm[:30]}'),
               use_numbers=match_mode.startswith('파일번호'))
    bar.empty()
    if owner_files:
        assign_owners(rows, read_owner_master([(f.name, f.getvalue()) for f in owner_files]))
    for r in rows:
        r['owner'] = r.get('owner') or '미지정'
    post = None
    if post_file:
        try:
            post, psh, pcols, used_qty = read_postcost(post_file.getvalue(), post_sheet, post_unit)
            st.session_state['post_info'] = (psh, pcols, used_qty)
        except Exception as e:
            st.warning(f'사후원가 파일을 읽지 못했습니다: {e}')
    st.session_state.update(rows=rows, files=files, post=post, cache={}, ver=0, saved_msg=None)

rows = st.session_state.get('rows')
if not rows:
    st.stop()

files = st.session_state['files']
path_by_name = {os.path.basename(p): p for p in files}

n = len(rows)
m = st.columns(5)
m[0].metric('요청 품목', n)
m[1].metric('✅ 완료(자동+확인)', sum(1 for r in rows if r['status'] == 'O' and not r['review']))
m[2].metric('🟡 검토 필요', sum(1 for r in rows if r['review']))
m[3].metric('⚪ 견적서 없음', sum(1 for r in rows if r['status'] == 'X' and not r.get('checked')))
m[4].metric('✔ 확인완료 체크', sum(1 for r in rows if r.get('checked')))

tab1, tab2, tab3 = st.tabs(['① 자동입력 결과 · 수정/저장', '② 담당자별 요청 메일', '③ 사전 vs 사후원가'])

NUM = ['원재료비', '부원료비', '부재료비', '가공비', '마진', '최종견적']
KEY = {'원재료비': 'raw', '부원료비': 'sub', '부재료비': 'pack', '가공비': 'proc', '마진': 'margin', '최종견적': 'final'}

with tab1:
    if st.session_state.get('saved_msg'):
        st.success(st.session_state['saved_msg'])
    df = pd.DataFrame([{
        '확인': bool(r.get('checked')),
        '구분': ('✔' if r.get('checked') else '⚪' if r['status'] == 'X' else '🟡' if r['review'] else '✅'),
        'no': r['no'], '제품명': r['name'], '거래처': r['customer'], '담당자': r.get('owner'), '단가': r['price'],
        '원재료비': r['raw'], '부원료비': r['sub'], '부재료비': r['pack'], '가공비': r['proc'],
        '마진': r['margin'], '최종견적': r['final'], 'O/X': r['status'],
        '매칭 견적서': r['file'] or '',
        '메모': ' / '.join(dict.fromkeys(r.get('edits', []) + r['notes'])),
    } for r in rows])

    f1, f2 = st.columns([1, 3])
    only = f1.toggle('검토 필요·견적서 없음만 보기', value=False)
    f2.caption('① 숫자·담당자·O/X를 고치거나 ② **매칭 견적서**를 다른 파일로 바꾸고(그 파일로 다시 읽음) '
               '③ 맞으면 **확인** 체크 → 반드시 **💾 저장**을 눌러야 아래 엑셀 다운로드에 반영됩니다.')
    view = df[df['구분'].isin(['🟡', '⚪'])] if only else df

    with st.form(f"edit_{st.session_state['ver']}"):
        edited = st.data_editor(
            view, use_container_width=True, height=540, hide_index=True,
            disabled=['구분', 'no', '제품명', '거래처', '단가', '메모'],
            column_config={
                '확인': st.column_config.CheckboxColumn('확인', help='내용을 확인했으면 체크 → 노란색 해제', width='small'),
                'O/X': st.column_config.SelectboxColumn('O/X', options=['O', 'X'], width='small'),
                '매칭 견적서': st.column_config.SelectboxColumn(
                    '매칭 견적서', options=[''] + sorted(path_by_name), width='large',
                    help='다른 견적서를 고르면 저장 시 그 파일로 다시 추출합니다'),
                **{k: st.column_config.NumberColumn(k, format='%d', step=1) for k in ['단가'] + NUM},
            })
        save = st.form_submit_button('💾 저장 (수정·확인 내용 반영)', type='primary')

    if save:
        cache = st.session_state['cache']
        n_edit = n_file = n_chk = 0
        for idx, er in edited.iterrows():
            r = rows[idx]
            orig = df.loc[idx]
            # 1) 견적서 파일 변경 → 재추출
            new_file = er['매칭 견적서'] or ''
            if new_file != (orig['매칭 견적서'] or ''):
                r = reselect_file(r, path_by_name.get(new_file), cache)
                rows[idx] = r
                n_file += 1
            else:
                # 2) 숫자 수정 (원래 값과 다를 때만, 변경 이력 기록)
                for col in NUM:
                    nv, ov = er[col], orig[col]
                    nv = None if pd.isna(nv) else int(nv)
                    ov = None if pd.isna(ov) else int(ov)
                    if nv != ov:
                        r[KEY[col]] = nv
                        r.setdefault('edits', []).append(f"{col} {ov if ov is not None else '-'}→{nv if nv is not None else '-'}")
                        n_edit += 1
                if er['O/X'] != orig['O/X']:
                    r['status'] = er['O/X']
                    r.setdefault('edits', []).append(f"O/X {orig['O/X']}→{er['O/X']}")
                    n_edit += 1
            if (er['담당자'] or '') != (orig['담당자'] or ''):
                r['owner'] = er['담당자'] or '미지정'
                r['owner_src'] = '수동'
                n_edit += 1
            if bool(er['확인']) != bool(orig['확인']):
                r['checked'] = bool(er['확인'])
                n_chk += 1
            refresh_review(r)
        st.session_state['ver'] += 1
        st.session_state['saved_msg'] = (f'저장 완료 — 값 수정 {n_edit}건 · 견적서 변경 {n_file}건 · 확인 체크 변경 {n_chk}건. '
                                         f'엑셀 다운로드에 반영되었습니다.')
        st.rerun()

with tab2:
    reqs = owner_requests(rows)
    if not reqs:
        st.success('요청할 품목이 없습니다.')
    st.caption('확인 체크한 품목은 요청 대상에서 빠집니다. 이메일을 넣으면 [메일 쓰기]로 메일 프로그램이 열립니다(자동 발송 없음).')
    emails = st.session_state.setdefault('emails', {})
    for owner, lst in reqs.items():
        subj, body = request_mail(owner, lst, period, sender, due or None)
        with st.expander(f'{owner} — {len(lst)}건', expanded=False):
            emails[owner] = st.text_input('이메일', value=emails.get(owner, ''), key=f'em_{owner}')
            st.text_input('제목', value=subj, key=f'sj_{owner}')
            st.text_area('본문', value=body, height=260, key=f'bd_{owner}')
            if emails[owner]:
                url = f"mailto:{emails[owner]}?subject={urllib.parse.quote(subj)}&body={urllib.parse.quote(body)}"
                st.link_button('✉️ 메일 쓰기', url)

post = st.session_state.get('post')
cmp = compare_pre_post(rows, post) if post else None
with tab3:
    if not cmp:
        st.info('④ 사후원가 파일을 올리면 품목별 사전(견적) vs 사후(실적) 단위원가 차이를 계산합니다.')
    else:
        psh, pcols, used_qty = st.session_state['post_info']
        st.caption(f"시트 '{psh}' · 생산량 나눔 {'예' if used_qty else '아니오'} · "
                   f"원재료={pcols['raw']} / 부재료={pcols['pack']} / 가공비={pcols['proc']}")
        cdf = pd.DataFrame([{
            'no': c['no'], '제품명': c['name'], '거래처': c['customer'], '담당자': c.get('owner'),
            '최종견적': c['final'], '사전원가': c['pre_cost'], '사후원가': c.get('post_cost'),
            '차이': c.get('diff'), '차이율%': None if c.get('diff_rate') is None else c['diff_rate'] * 100,
            '사전이익률%': None if c.get('pre_margin_rate') is None else c['pre_margin_rate'] * 100,
            '사후이익률%': None if c.get('post_margin_rate') is None else c['post_margin_rate'] * 100,
            '매칭': c['match']} for c in cmp])
        k = st.columns(3)
        k[0].metric('사후원가 매칭', f"{(cdf['사후원가'].notna()).sum()} / {len(cdf)}")
        k[1].metric('원가 5% 초과 상승', int((cdf['차이율%'] > 5).sum()))
        k[2].metric('사후 적자 품목', int((cdf['사후이익률%'] < 0).sum()))
        st.dataframe(cdf.sort_values('차이율%', ascending=False, na_position='last'), use_container_width=True,
                     hide_index=True, height=480)

st.divider()
title = f'{period} 사전원가 요청 품목 목록 (자동입력)'
st.caption('💾 저장한 내용까지 반영됩니다. 표에서 고친 뒤 저장을 누르지 않으면 반영되지 않습니다.')
st.download_button('📥 엑셀 다운로드 (->영업요청 · 담당자별 요청 · 사전vs사후)',
                   data=to_excel(rows, title, compare=cmp, period=period),
                   file_name=f'{title}.xlsx', type='primary',
                   mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
