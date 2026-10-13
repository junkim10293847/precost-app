"""사전원가 자동입력 핵심 로직 (UI 독립).

1) 요청 품목 목록(엑셀) 읽기
2) 견적서 파일들 ↔ 품목 매칭 (파일번호 + 품목명 유사도)
3) 견적서에서 원가 항목 추출
4) '->영업요청' 형식 엑셀로 출력
"""
import io, os, re, zipfile, tempfile, difflib, warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from extractor import extract_file, file_numbers

warnings.filterwarnings('ignore')

SUPPORTED = ('.xlsx', '.xlsm', '.xls', '.pdf', '.png', '.jpg', '.jpeg')

# ---------------------------------------------------------------- 목록 읽기
HEADER_SYNONYMS = {
    'no': ['no', 'No', 'NO', '번호'],
    'name': ['제품명', '품목명'],
    'lot': ['LOT', 'lot', '제조번호', '로트'],
    'customer': ['거래처', '고객사'],
    'price': ['단가'],
    'spec': ['규격'],
    'owner': ['담당자'],
}


def read_item_list(path_or_bytes, sheet=None):
    wb = openpyxl.load_workbook(path_or_bytes if isinstance(path_or_bytes, str) else io.BytesIO(path_or_bytes),
                                data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    hdr_row, cols = None, {}
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = {c: str(ws.cell(r, c).value).strip() for c in range(1, ws.max_column + 1)
                if ws.cell(r, c).value is not None}
        if '제품명' in vals.values() or '품목명' in vals.values():
            hdr_row = r
            for c, v in vals.items():
                for key, syns in HEADER_SYNONYMS.items():
                    if v in syns:
                        cols.setdefault(key, []).append(c)
            break
    if hdr_row is None:
        raise ValueError('목록에서 "제품명" 헤더를 찾지 못했습니다.')

    def pick(key, r):
        best = None
        for c in cols.get(key, []):
            v = ws.cell(r, c).value
            if v is None or (isinstance(v, str) and v.startswith('#')):
                continue
            best = v  # 같은 헤더가 여러 개면 값이 있는 마지막 열
        return best

    items = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        name = pick('name', r)
        if not name or not str(name).strip():
            continue
        items.append(dict(row=r, no=pick('no', r), name=str(name).strip(), lot=pick('lot', r),
                          customer=pick('customer', r), price=pick('price', r), spec=pick('spec', r),
                          owner=pick('owner', r), tag=ws.cell(r, 1).value))
    return items, ws.title


# ---------------------------------------------------------------- 파일 모으기
def _decode_zip_name(info):
    n = info.filename
    if not (info.flag_bits & 0x800):
        try:
            n = n.encode('cp437').decode('cp949')
        except Exception:
            pass
    return n


def collect_files(uploads, workdir):
    """uploads: [(filename, bytes)] — zip은 풀어서 전개."""
    out = []
    for fname, data in uploads:
        if fname.lower().endswith('.zip'):
            z = zipfile.ZipFile(io.BytesIO(data))
            for info in z.infolist():
                n = os.path.basename(_decode_zip_name(info))
                if not n or info.is_dir() or n.startswith('~$') or not n.lower().endswith(SUPPORTED):
                    continue
                p = os.path.join(workdir, n)
                with open(p, 'wb') as f:
                    f.write(z.read(info))
                out.append(p)
        elif fname.lower().endswith(SUPPORTED):
            p = os.path.join(workdir, os.path.basename(fname))
            with open(p, 'wb') as f:
                f.write(data)
            out.append(p)
    return sorted(set(out))


# ---------------------------------------------------------------- 매칭
def _clean(s):
    s = re.sub(r'^\s*[\d~.\s]+\.', '', s)                       # 앞 번호
    s = re.sub(r'\.(xlsx|xls|xlsm|pdf|png|jpe?g)$', '', s, flags=re.I)
    s = re.sub(r'20\d{2}\s*\.\s*\d{1,2}\s*\.\s*\d{1,2}|\b\d{6,8}\b', ' ', s)  # 날짜
    s = re.sub(r'견적서|사전원가계산서|확\s*정|최종|견본일지|복사본|사본|수정', ' ', s)
    return re.sub(r'[\s\(\)\[\]_\-·∙,&+]', '', s).lower()


def _nums(s):
    return set(re.findall(r'\d{3,}', re.sub(r'20\d{2}[.\s]*\d{1,2}[.\s]*\d{1,2}|\d{6,8}', '', s.replace(',', ''))))


def spec_counts(s):
    return {int(x) for x in re.findall(r'(\d+)\s*(?:포|정|T|캡슐|캅셀|입|개입|병)', s.replace(' ', ''))}


def name_score(item_name, fname, spec=None):
    a, b = _clean(item_name), _clean(os.path.basename(fname))
    base = difflib.SequenceMatcher(None, a, b).ratio()
    # 부분 포함 보정
    sub = difflib.SequenceMatcher(None, a, b).find_longest_match(0, len(a), 0, len(b)).size / max(len(a), 1)
    sc = max(base, sub * 0.9)
    na, nb = _nums(item_name), _nums(os.path.basename(fname))
    if na and nb:
        sc += 0.15 if na & nb else -0.15   # 1200 vs 1500 같은 함량 숫자 일치 여부
    if isinstance(spec, (int, float)) and spec and spec_counts(os.path.basename(fname)):
        sc += 0.1 if int(spec) in spec_counts(os.path.basename(fname)) else -0.3   # 30포 vs 90포
    return round(sc, 3)


def match_items(items, files, use_numbers=True, min_name=0.5):
    by_no = {}
    for f in (files if use_numbers else []):
        for n in file_numbers(f):
            by_no.setdefault(n, []).append(f)
    item_by_no = {it['no']: it for it in items if isinstance(it['no'], int)}

    def owner_score(f):
        """파일번호 기준 '원래 주인' 품목이 이 파일과 얼마나 맞는지"""
        scs = [name_score(item_by_no[n]['name'], f) for n in file_numbers(f) if n in item_by_no]
        return max(scs) if scs else 0

    results = []
    for it in items:
        no = it['no'] if isinstance(it['no'], int) else None
        scored = sorted(((name_score(it['name'], f, it.get('spec')), f) for f in files), reverse=True)
        best_sc, best_f = scored[0] if scored else (0, None)
        num_cands = sorted(((name_score(it['name'], f, it.get('spec')), f) for f in by_no.get(no, [])), reverse=True)
        steal_ok = best_f is not None and best_sc >= 0.45 and owner_score(best_f) < best_sc - 0.1
        if not use_numbers:
            if best_f and best_sc >= min_name:
                results.append((best_f, best_sc, '품목명매칭' + ('' if best_sc >= 0.75 else '-확인필요')))
            else:
                results.append((None, 0, '견적서없음'))
            continue
        if num_cands:
            sc, f = num_cands[0]
            if best_f != f and steal_ok and best_sc - sc >= 0.2:
                results.append((best_f, best_sc, '번호불일치→품목명으로 매칭'))
            else:
                sc0 = name_score(it['name'], f)   # 규격 보정 전 점수
                results.append((f, sc, '번호매칭' + ('' if sc0 >= 0.3 else ' (품목명 상이-확인필요)')))
        elif steal_ok:
            results.append((best_f, best_sc, '품목명매칭(해당 번호 파일없음)-확인필요'))
        else:
            results.append((None, 0, '견적서없음'))
    return results


# ---------------------------------------------------------------- xls 지원 (서버에 xlrd 없을 때 LibreOffice 변환)
def ensure_readable(path):
    if not path.lower().endswith('.xls'):
        return path
    try:
        import xlrd  # noqa
        return path
    except ImportError:
        import subprocess
        outdir = os.path.join(os.path.dirname(path), '_conv')
        os.makedirs(outdir, exist_ok=True)
        subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', outdir, path],
                       capture_output=True, timeout=120)
        cand = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + '.xlsx')
        return cand if os.path.exists(cand) else path


# ---------------------------------------------------------------- 실행
def fill_row(it, f, how='수동선택', sc=1.0, cache=None):
    """품목 1행 + 견적서 파일 1개 → 결과 레코드"""
    cache = {} if cache is None else cache
    rec = dict(it)
    rec.update(file=os.path.basename(f) if f else '', file_path=f, match=how, match_score=sc, ocr=False,
               raw=None, sub=None, pack=None, proc=None, margin=None, final=None, computed=None, notes=[],
               checked=False, edits=[])
    if f:
        if f not in cache:
            try:
                cache[f] = extract_file(ensure_readable(f))
            except Exception as e:  # noqa
                cache[f] = ({}, {}, None, False)
                rec['notes'].append(f'읽기오류: {str(e)[:40]}')
        res, src, sheet, ocr = cache[f]
        rec['ocr'] = ocr
        rec['sheet'] = sheet
        for k in ('raw', 'pack', 'proc', 'margin', 'final'):
            v = res.get(k)
            rec[k] = round(v) if isinstance(v, (int, float)) else None
        rec['sub'] = None   # 부원료비는 사용하지 않음
        rec['computed'] = res.get('_computed_final')
        rec['src'] = src
        miss = [lbl for k, lbl in (('raw', '원재료비'), ('pack', '부재료비'), ('proc', '가공비'), ('final', '최종견적'))
                if rec[k] is None]
        if miss:
            rec['notes'].append('미추출:' + '/'.join(miss))
        if ocr:
            rec['notes'].append('스캔/이미지(OCR) - 확인필요')
        fs = spec_counts(os.path.basename(f))
        if isinstance(it.get('spec'), (int, float)) and it['spec'] and fs and int(it['spec']) not in fs:
            rec['notes'].append(f"규격 상이-확인필요(목록 {int(it['spec'])} / 견적서 {'·'.join(map(str, sorted(fs)))})")
        if all(rec[k] is not None for k in ('raw', 'pack', 'proc', 'final')):
            base = rec['raw'] + rec['pack'] + rec['proc'] + (rec['margin'] or 0)
            if rec['sub'] and abs(rec['final'] - base) <= abs(rec['final'] - base - rec['sub']):
                rec['sub'] = None   # 요약표 옆 참고계산값 → 부원료비 아님
            gap = rec['final'] - base - (rec['sub'] or 0)
            if abs(gap) > max(30, rec['final'] * 0.02):
                rec['notes'].append(f'참고: 견적조정 {gap:+,.0f}')
        if isinstance(it.get('price'), (int, float)) and rec['final'] and abs(it['price'] - rec['final']) >= 1:
            rec['notes'].append(f"참고: 단가-견적 {it['price'] - rec['final']:+,.0f}")
    rec['status'] = 'O' if f else 'X'
    if not f:
        rec['notes'].append('견적서없음')
    if '확인' in how or '불일치' in how:
        rec['notes'].append(how)
    refresh_review(rec)
    return rec


def refresh_review(rec):
    rec['review'] = (not rec.get('checked')) and rec['status'] == 'O' and any(
        k in ' '.join(rec['notes']) for k in ('미추출', 'OCR', '불일치', '확인필요', '오류'))
    return rec


def run(items, files, progress=None, use_numbers=True):
    matches = match_items(items, files, use_numbers=use_numbers)
    cache, rows = {}, []
    for i, (it, (f, sc, how)) in enumerate(zip(items, matches)):
        if progress:
            progress(i + 1, len(items), it['name'])
        rows.append(fill_row(it, f, how, sc, cache))
    return rows


ITEM_KEYS = ('row', 'no', 'name', 'lot', 'customer', 'price', 'spec', 'owner', 'tag')


def reselect_file(rec, path, cache=None):
    """사용자가 견적서 파일을 다시 고른 경우 해당 파일로 재추출 (담당자·확인 상태는 유지)"""
    it = {k: rec.get(k) for k in ITEM_KEYS}
    new = fill_row(it, path, '수동선택' if path else '견적서없음', 1.0, cache)
    new['owner_src'] = rec.get('owner_src', '')
    new['edits'] = rec.get('edits', []) + [f"견적서 변경: {rec.get('file') or '없음'} → {new['file'] or '없음'}"]
    if isinstance(it.get('price'), (int, float)) and it['price'] and new['final'] and \
            abs(it['price'] - new['final']) / it['price'] > 0.1:
        new['notes'].append('수동선택 견적서의 견적가가 단가와 10% 이상 차이-확인필요')
        refresh_review(new)
    return new


# ---------------------------------------------------------------- 엑셀 출력 (->영업요청 형식)
HEAD = ['no', '제품명', 'LOT', '거래처', '단가', '규격', '담당자', 'O/X', '비고', '원재료비',
        '부재료비', '가공비', '제조원가', '마진', '출고가', '최종견적', '+/-조정', '조정이익', '검증', '비고', 'F-Q',
        '매칭 견적서 파일', '매칭방식', '자동검토 메모']


def to_excel(rows, title='사전원가 요청 품목 목록', compare=None, period='', requests_sheet=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '->영업요청'
    thin = Side(style='thin', color='BFBFBF')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws['B2'] = title
    ws['B2'].font = Font(bold=True, size=14)
    ws['H3'] = '* 견적서 파일에서 자동 추출 (노란색 = 검토 필요, 회색 = 견적서 없음, 연두색 = 담당자 확인완료)'
    ws['H3'].font = Font(color='C00000', size=9)
    for j, h in enumerate(HEAD, start=2):
        c = ws.cell(4, j, h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F4E78')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = bd
    yellow = PatternFill('solid', fgColor='FFF2CC')
    grey = PatternFill('solid', fgColor='EDEDED')
    green = PatternFill('solid', fgColor='E2EFDA')
    for i, r in enumerate(rows):
        n = 5 + i
        L = lambda col: f'{col}{n}'  # noqa
        ws.cell(n, 1, r.get('tag'))
        vals = [r['no'], r['name'], r['lot'], r['customer'], r['price'], r['spec'], (r.get('owner') if r.get('owner') != '미지정' else None), r['status'],
                None, r['raw'], r['pack'], r['proc'], f'=SUM({L("K")}:{L("M")})', r['margin'],
                f'=+{L("N")}+{L("O")}', r['final'] if r['status'] == 'O' else f'=+{L("F")}',
                f'=+{L("P")}-{L("Q")}', f'=+{L("Q")}-{L("N")}', f'=+{L("N")}+{L("O")}-{L("R")}={L("Q")}',
                '견적서없음' if r['status'] == 'X' else None, f'=+{L("F")}-{L("Q")}',
                r['file'], r['match'],
                ' / '.join(dict.fromkeys((['✔확인완료'] if r.get('checked') else []) + r.get('edits', []) + r['notes']))]
        for j, v in enumerate(vals, start=2):
            c = ws.cell(n, j, v)
            c.border = bd
            if j in (6, 11, 12, 13, 14, 15, 16, 17, 18, 19, 22):
                c.number_format = '#,##0'
            if r['status'] == 'X' and not r.get('checked'):
                c.fill = grey
            elif r.get('checked'):
                c.fill = green
            elif r['review']:
                c.fill = yellow
    widths = dict(A=8, B=5, C=38, D=11, E=16, F=9, G=6, H=8, I=5, J=8, K=9, L=9, M=9, N=9, O=8, P=9, Q=9,
                  R=8, S=8, T=8, U=10, V=8, W=42, X=22, Y=44)
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    ws.freeze_panes = 'D5'
    ws.auto_filter.ref = f'B4:Y{4 + len(rows)}'
    hdr_style = lambda c: (setattr(c, 'font', Font(bold=True, color='FFFFFF')),  # noqa
                           setattr(c, 'fill', PatternFill('solid', fgColor='1F4E78')),
                           setattr(c, 'alignment', Alignment(horizontal='center', wrap_text=True)))

    # ---- 담당자별 요청
    reqs = owner_requests(rows)
    w2 = wb.create_sheet('담당자별 요청')
    w2['A1'] = f'{period} 사전원가 — 담당자별 견적서 요청 목록'.strip()
    w2['A1'].font = Font(bold=True, size=13)
    h2 = ['담당자', 'no', '제품명', '거래처', 'LOT', '요청 사유', '담당자 근거']
    for j, h in enumerate(h2, 1):
        hdr_style(w2.cell(3, j, h))
    rr = 4
    for owner, lst in reqs.items():
        for r, why in lst:
            for j, v in enumerate([owner, r['no'], r['name'], r['customer'], r['lot'], why, r.get('owner_src', '')], 1):
                w2.cell(rr, j, v).border = bd
            rr += 1
    for k, w in dict(A=10, B=6, C=44, D=18, E=12, F=34, G=14).items():
        w2.column_dimensions[k].width = w
    w2.freeze_panes = 'A4'

    # ---- 사전 vs 사후
    if compare:
        w3 = wb.create_sheet('사전vs사후')
        w3['A1'] = f'{period} 사전원가(견적) vs 사후원가(실적) — 단위당 원'.strip()
        w3['A1'].font = Font(bold=True, size=13)
        h3 = ['no', '제품명', '거래처', 'LOT', '담당자', '최종견적', '사전 원재료', '사후 원재료', '사전 부재료',
              '사후 부재료', '사전 가공비', '사후 가공비', '사전 원가계', '사후 원가계', '차이(사후-사전)', '차이율',
              '사전 이익률', '사후 이익률', '사후원가 매칭']
        for j, h in enumerate(h3, 1):
            hdr_style(w3.cell(3, j, h))
        red = PatternFill('solid', fgColor='F8CBAD')
        for i, c in enumerate(compare):
            n = 4 + i
            vals = [c['no'], c['name'], c['customer'], c['lot'], c.get('owner'), c['final'], c['pre_raw'],
                    c.get('post_raw'), c['pre_pack'], c.get('post_pack'), c['pre_proc'], c.get('post_proc'),
                    c['pre_cost'], c.get('post_cost'), c.get('diff'), c.get('diff_rate'),
                    c.get('pre_margin_rate'), c.get('post_margin_rate'), c['match']]
            for j, v in enumerate(vals, 1):
                cell = w3.cell(n, j, round(v) if isinstance(v, float) and j < 16 else v)
                cell.border = bd
                if 6 <= j <= 15:
                    cell.number_format = '#,##0'
                if j >= 16 and j <= 18:
                    cell.number_format = '0.0%'
            if (c.get('diff_rate') or 0) > 0.05 or (c.get('post_margin_rate') is not None and c['post_margin_rate'] < 0):
                for j in range(1, len(h3) + 1):
                    w3.cell(n, j).fill = red
        for k, w in dict(A=6, B=36, C=16, D=11, E=8).items():
            w3.column_dimensions[k].width = w
        w3.freeze_panes = 'C4'
    if not requests_sheet:
        del wb['담당자별 요청']
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ================================================================ Q3. 담당자 배정 · 요청 메일
def _norm_key(s):
    return re.sub(r'[\s㈜\(\)주식회사]', '', str(s or '')).lower()


def read_owner_master(sources):
    """이전 결과 파일(->영업요청 / ->세부사항 등)에서 거래처→담당자, 제품명→담당자 사전 생성.
    sources: [(파일명, bytes)]"""
    import collections
    by_cust, by_name = collections.defaultdict(collections.Counter), {}
    for fname, data in sources:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 2000), values_only=True))
            for hi, row in enumerate(rows[:15]):
                hdr = [str(v).strip() if v is not None else '' for v in row]
                if '담당자' in hdr and '제품명' in hdr and '거래처' in hdr:
                    io_, ic, iw = hdr.index('담당자'), hdr.index('제품명'), hdr.index('거래처')
                    for r in rows[hi + 1:]:
                        if r is None or len(r) <= max(io_, ic, iw):
                            continue
                        o, n, c = r[io_], r[ic], r[iw]
                        if o and n:
                            by_name[_norm_key(n)] = str(o).strip()
                            if c:
                                by_cust[_norm_key(c)][str(o).strip()] += 1
                    break
    return {'cust': {k: v.most_common(1)[0][0] for k, v in by_cust.items()},
            'cust_multi': {k: list(v) for k, v in by_cust.items() if len(v) > 1},
            'name': by_name}


def assign_owners(rows, master):
    for r in rows:
        if r.get('owner'):
            r['owner_src'] = '목록'
            continue
        nk, ck = _norm_key(r['name']), _norm_key(r['customer'])
        if nk in master['name']:
            r['owner'], r['owner_src'] = master['name'][nk], '제품명'
        elif ck in master['cust']:
            r['owner'] = master['cust'][ck]
            r['owner_src'] = '거래처' + ('(복수담당-확인)' if ck in master['cust_multi'] else '')
        else:
            r['owner'], r['owner_src'] = '미지정', ''
    return rows


def needs_request(r):
    """영업 담당자에게 요청해야 할 행: 견적서 없음 / 값 미추출 / 스캔본 / 매칭 의심"""
    if r.get('checked'):
        return None
    if r['status'] == 'X':
        return '견적서 미제출'
    t = ' '.join(r['notes'])
    if '미추출' in t or 'OCR' in t:
        return '견적서 판독 불가(값 확인 요청)'
    if '규격 상이' in t:
        return '규격 확인 요청(' + t.split('규격 상이-확인필요(')[1].split(')')[0] + ')'
    if '불일치' in t or '확인' in t:
        return '견적서 파일 확인 요청(품목명/번호 상이)'
    return None


def owner_requests(rows):
    out = {}
    for r in rows:
        why = needs_request(r)
        if why:
            out.setdefault(r.get('owner') or '미지정', []).append((r, why))
    return dict(sorted(out.items(), key=lambda kv: (kv[0] == '미지정', kv[0])))


def request_mail(owner, reqs, period='이번 달', sender='재무지원팀', due=None):
    lines = [f'{owner}님, 안녕하세요. {sender}입니다.', '',
             f'{period} 사전원가 집계 중 아래 {len(reqs)}개 품목의 견적서 확인이 필요합니다.',
             '견적서 파일을 회신해 주시거나 내용 확인 부탁드립니다.' + (f' (회신 기한: {due})' if due else ''), '']
    for r, why in reqs:
        lines.append(f"- [{r['no']}] {r['name']} / {r['customer'] or ''} / LOT {r['lot'] or ''} → {why}")
    lines += ['', '※ 견적서 파일명 맨 앞에 목록 번호를 붙여 주세요. (예: 29. 테아닌 견적서.xlsx)', '', '감사합니다.']
    subject = f'[사전원가] {period} 견적서 요청 - {owner}님 {len(reqs)}건'
    return subject, '\n'.join(lines)


# ================================================================ Q2. 사전원가 vs 사후원가
POST_MAP = {
    'raw': ['원료비', '원재료비'],
    'pack': ['부자재비', '부재료비', '자재비'],
    'proc': ['직접노무비', '제조경비(간접)', '간접노무비', '제조경비(직접)+감가상각비', '외주비용', '세현포장비',
             '가공비', '노무비', '제조경비', '감가상각비'],
    'qty': ['생산량', '생산수량'],
    'lot': ['제조번호', 'LOT', 'lot', '로트'],
    'name': ['제품명', '품목명'],
}


def _val(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(',', '')
        try:
            return float(s)
        except ValueError:
            return 0.0  # ' - ', '#DIV/0!' 등
    return 0.0


def read_postcost(data, sheet=None, per_unit=None, mapping=None):
    """사후원가 엑셀 → {LOT키: {...}, '__name__'+제품명키: {...}}
    per_unit=None 이면 '생산량' 열이 있으면 금액÷생산량, 없으면 이미 단위원가로 간주."""
    mp = mapping or POST_MAP
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheets = [wb[sheet]] if sheet else wb.worksheets
    for ws in sheets:
        for hr in range(1, min(ws.max_row, 15) + 1):
            hdr = {str(ws.cell(hr, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value}
            if not any(h in hdr for h in mp['name']) or not any(h in hdr for h in mp['raw']):
                continue
            col = {k: [hdr[h] for h in hs if h in hdr] for k, hs in mp.items()}
            use_qty = (per_unit is None and bool(col['qty'])) or per_unit is False
            out = {}
            for r in range(hr + 1, ws.max_row + 1):
                name = ws.cell(r, col['name'][0]).value if col['name'] else None
                if not name:
                    continue
                rec = {k: sum(_val(ws.cell(r, c).value) for c in col[k]) for k in ('raw', 'pack', 'proc')}
                qty = _val(ws.cell(r, col['qty'][0]).value) if col['qty'] else 0
                if use_qty:
                    if qty <= 0:
                        continue
                    rec = {k: v / qty for k, v in rec.items()}
                rec.update(name=str(name), qty=qty, lot=ws.cell(r, col['lot'][0]).value if col['lot'] else None,
                           has=any(rec[k] for k in ('raw', 'pack', 'proc')))
                if rec['lot'] is not None:
                    out[_norm_key(rec['lot'])] = rec
                out['__name__' + _norm_key(name)] = rec
            return out, ws.title, {k: [ws.cell(hr, c).value for c in v] for k, v in col.items()}, use_qty
    raise ValueError('사후원가 파일에서 "제품명"과 "원료비" 머리글을 찾지 못했습니다.')


def compare_pre_post(rows, post):
    res = []
    for r in rows:
        p = post.get(_norm_key(r['lot'])) if r.get('lot') is not None else None
        how = 'LOT'
        if p is None:
            p, how = post.get('__name__' + _norm_key(r['name'])), '제품명'
        pre_cost = sum(r[k] or 0 for k in ('raw', 'pack', 'proc')) if r['status'] == 'O' else None
        item = dict(no=r['no'], name=r['name'], customer=r['customer'], lot=r['lot'], owner=r.get('owner'),
                    final=r['final'], pre_raw=r['raw'], pre_pack=r['pack'], pre_proc=r['proc'], pre_cost=pre_cost,
                    match=how if p else '사후원가 없음')
        if p and p['has']:
            post_cost = p['raw'] + p['pack'] + p['proc']
            item.update(post_raw=p['raw'], post_pack=p['pack'], post_proc=p['proc'], post_cost=post_cost)
            if pre_cost:
                item['diff'] = post_cost - pre_cost
                item['diff_rate'] = item['diff'] / pre_cost
            price = r['price'] or r['final']
            if price:
                item['pre_margin_rate'] = (price - pre_cost) / price if pre_cost else None
                item['post_margin_rate'] = (price - post_cost) / price
        res.append(item)
    return res


# ================================================================ 견적서 미리보기 (검토 화면용)
def preview_quote(path, sheet=None, max_pages=3, resolution=110):
    """견적서 원본 미리보기.
    반환: ('images', [PIL.Image...]) | ('tables', [(시트명, DataFrame)...]) | ('none', 메시지)"""
    import pandas as pd
    if not path or not os.path.exists(path):
        return 'none', '견적서 파일이 없습니다.'
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == '.pdf':
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return 'images', [p.to_image(resolution=resolution).original for p in pdf.pages[:max_pages]]
        if ext in ('.png', '.jpg', '.jpeg'):
            from PIL import Image
            return 'images', [Image.open(path)]
        if ext in ('.xlsx', '.xlsm', '.xls'):
            p = ensure_readable(path)
            out = []
            if p.lower().endswith('.xls'):
                import xlrd
                wb = xlrd.open_workbook(p)
                sheets = [(sh.name, [sh.row_values(r) for r in range(sh.nrows)]) for sh in wb.sheets()]
            else:
                wb = openpyxl.load_workbook(p, data_only=True)
                sheets = [(ws.title, [[c for c in row] for row in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
            for title, data in sheets:
                if not data:
                    continue
                df = pd.DataFrame(data[:150])
                df = df.replace('', None).dropna(how='all').dropna(axis=1, how='all')
                if df.empty:
                    continue
                df = df.map(lambda v: (round(v) if abs(v) >= 100 else round(v, 2)) if isinstance(v, float) else v)
                df.columns = [openpyxl.utils.get_column_letter(int(c) + 1) for c in df.columns]
                df.index = [int(i) + 1 for i in df.index]
                out.append((title, df.astype(str).replace({'None': '', 'nan': ''})))
            # 추출에 쓴 시트를 맨 앞으로
            out.sort(key=lambda t: (t[0] != sheet, '견본' in t[0]))
            return 'tables', out
    except Exception as e:  # noqa
        return 'none', f'미리보기 오류: {str(e)[:80]}'
    return 'none', '미리보기를 지원하지 않는 형식입니다.'
